"""엑셀 방식: 사용자가 가진 장부 엑셀(양식 제각각) → 정산 엑셀 하나.

구글 시트 방식과 같은 탭 구성으로 만든다. 그래서 마지막 단계(deliver.py)를 같이 쓴다.
  요약 · 부의금 입력 · 근조 화환 · 지출 · 물품 통계 · 확인 필요 · 사람별 탭(한 분당 하나)

- 탭마다 제목 줄(위쪽 10줄 안)을 보고 부의금 / 조화·물품 / 지출 중 무엇인지 알아본다.
- 받은 가족 열이 없으면 시트 이름(여러 시트일 때) → 파일 이름을 받은 가족으로 쓴다.
- 금액은 50000 · 50,000원 · ₩50,000 · 5만 · 10만5천 을 읽는다. 1,000 미만 숫자·글자는 추측하지 않고 '확인 필요'.
- 이름 칸에 '소계'가 든 줄은 일차 구분선, '합계'·'총계'는 건너뛴다. 소계 값은 '장부에 적은 소계'로 검산한다.
- 원본 엑셀은 읽기만 한다. 결과는 계산된 값으로 저장한다(입력이 바뀌면 다시 돌린다).

    python3 build_excel.py 장부.xlsx [더.xlsx ...] -o 부의금정산.xlsx [--split 1|2|3] [--exclude 이름,이름]

--split 은 사용자가 정산 방법을 고른 뒤에만 준다(1 똑같이 나누기, 2 받은 부의금 비율, 3 손님 수 비율).
고르기 전에는 split_preview.py 로 세 방법 결과를 보여 주고 묻는다.
"""
import argparse
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from split_preview import METHODS, shares

MONEY_KEYS = {
    "giver": ["입금자", "보낸분", "보낸사람", "보내는분", "조문객", "성명", "성함", "이름", "name"],
    "amount": ["입금액", "부의금", "조의금", "금액", "액수", "amount"],
}
ITEM_KEYS = {
    "giver": ["보낸분", "보낸곳", "보낸사람", "보내는분", "발신", "성명", "이름"],
    "kind": ["종류", "품목", "구분", "물품"],
}
SPEND_KEYS = {
    "item": ["항목", "내역", "지출", "사용처", "용도"],
    "amount": ["금액", "지출액", "비용", "액수"],
    "payer": ["낸사람", "낸분", "결제", "지불", "계산"],
}
COMMON_KEYS = {
    "family": ["받은가족", "받은사람", "받는사람", "누구손님", "상주", "가족", "담당", "손님"],
    "org": ["소속", "관계", "회사", "직장"],
    "memo": ["비고", "메모", "참고", "특이사항"],
}
WON = '#,##0"원"'
GREY = PatternFill("solid", fgColor="EDEDED")
BLUE = PatternFill("solid", fgColor="DCE6F7")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
LINE = Side(style="thin", color="D9D9D9")
BOX = Border(top=LINE, bottom=LINE, left=LINE, right=LINE)


def norm(s):
    return re.sub(r"[\s·./()_\-]", "", str(s or "")).lower()


def text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def parse_amount(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        n = int(round(v))
        return n if n >= 1000 else None
    s = re.sub(r"[\s,₩원]", "", str(v))
    if re.fullmatch(r"\d+", s):
        n = int(s)
        return n if n >= 1000 else None
    m = re.fullmatch(r"(?:(\d+(?:\.\d+)?)만)?(?:(\d+)천)?", s)
    if m and (m.group(1) or m.group(2)):
        return int(round(float(m.group(1) or 0) * 10000)) + int(m.group(2) or 0) * 1000
    return None


def match(row, keys):
    cols = {}
    for i, cell in enumerate(row):
        h = norm(cell)
        if not h:
            continue
        for field, words in list(keys.items()) + list(COMMON_KEYS.items()):
            if field not in cols and any(w in h for w in words):
                cols[field] = i
                break
    return cols


def classify(rows):
    """(종류, 제목 줄 번호, 열) - 지출 → 조화·물품 → 부의금 순으로 본다."""
    for i, row in enumerate(rows[:10]):
        c = match(row, SPEND_KEYS)
        if {"item", "amount"} <= c.keys() and "giver" not in match(row, MONEY_KEYS):
            return "spend", i, c
        c = match(row, ITEM_KEYS)
        if {"giver", "kind"} <= c.keys() and "amount" not in match(row, MONEY_KEYS):
            return "items", i, c
        c = match(row, MONEY_KEYS)
        if {"giver", "amount"} <= c.keys():
            return "money", i, c
    return None, None, None


def read_inputs(paths):
    money, items, spend = [], [], []
    ledger = defaultdict(int)  # 일차 → 장부에 적힌 소계(시트가 여럿이면 더한다)
    sheets_total = sum(len(load_workbook(p, read_only=True).sheetnames) for p in paths)
    for path in paths:
        wb = load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            kind, head, cols = classify(rows)
            if not kind:
                continue
            fallback = ws.title if sheets_total > 1 else Path(path).stem
            # 일차는 시트마다 따로 센다. 소계 줄이 없는 시트는 일차를 매기지 않는다.
            has_sub = kind == "money" and any("소계" in norm(row[cols["giver"]] if cols["giver"] < len(row) else "")
                                             for row in rows[head + 1:])
            day = 1 if has_sub else None
            for r, row in enumerate(rows[head + 1:], start=head + 2):
                get = lambda f: row[cols[f]] if f in cols and cols[f] < len(row) else None
                if all(text(c) == "" for c in row):
                    continue
                src = f"{Path(path).name} / {ws.title} / {r}행"
                fam = text(get("family")) if "family" in cols else fallback
                if kind == "money":
                    name = text(get("giver"))
                    if "소계" in norm(name):
                        v = parse_amount(get("amount"))
                        ledger[day] += v or 0
                        day += 1
                        continue
                    if norm(name) in ("합계", "총계", "총합계", "계", "total"):
                        continue
                    raw = get("amount")
                    money.append({"family": fam, "giver": name, "amount": parse_amount(raw), "raw": text(raw),
                                  "org": text(get("org")), "memo": text(get("memo")), "day": day, "src": src})
                elif kind == "items":
                    items.append({"family": fam, "giver": text(get("giver")), "kind": text(get("kind")),
                                  "org": text(get("org")), "memo": text(get("memo")), "src": src})
                else:
                    item = text(get("item"))
                    if re.search("합계|소계|총계", item):
                        continue
                    raw = get("amount")
                    spend.append({"item": item, "amount": parse_amount(raw), "raw": text(raw),
                                  "memo": text(get("memo")), "payer": text(get("payer")), "src": src})
    return money, items, spend, dict(ledger)


def checks(money, items, spend):
    key = lambda r: (r["family"], norm(r["giver"]), r["amount"])
    dup = Counter(key(r) for r in money if r["giver"] and r["amount"])
    for r in money:
        why = []
        if r["amount"] is None:
            why.append(f"금액 확인 ({r['raw'] or '빈칸'})")
        if not r["giver"]:
            why.append("이름 없음")
        if not r["family"]:
            why.append("받은 가족 없음")
        if r["giver"] and r["amount"] and dup[key(r)] > 1:
            why.append("중복 의심")
        r["check"] = ", ".join(why)
    # 같은 분이 종류를 달리해 여러 개 보내는 건 정상 - 보낸 분과 종류가 둘 다 같을 때만 의심
    idup = Counter((norm(r["giver"]), norm(r["kind"])) for r in items if r["giver"])
    for r in items:
        why = []
        if not r["family"]:
            why.append("받은 가족 없음")
        if not r["giver"]:
            why.append("보낸 분 없음")
        if r["giver"] and idup[(norm(r["giver"]), norm(r["kind"]))] > 1:
            why.append("중복 의심")
        r["check"] = ", ".join(why)
    for r in spend:
        why = []
        if r["amount"] is None:
            why.append(f"금액 확인 ({r['raw'] or '빈칸'})")
        if not r["payer"]:
            why.append("낸 사람 없음")
        r["check"] = ", ".join(why)


def head(ws, row, col, labels, fill=GREY):
    for i, h in enumerate(labels):
        c = ws.cell(row, col + i, h)
        c.font, c.fill, c.border = Font(bold=True), fill, BOX
        c.alignment = Alignment(horizontal="center")


def widths(ws, ws_widths):
    for col, w in ws_widths.items():
        ws.column_dimensions[col].width = w


def kpi(ws, row, col, label, value, fmt=None):
    ws.cell(row, col, label).font = Font(color="666666")
    c = ws.cell(row, col + 1, value)
    c.font = Font(bold=True, size=13)
    if fmt:
        c.number_format = fmt


def safe_title(name, used):
    t = re.sub(r"[\[\]:*?/\\]", "", name or "(미지정)")[:28] or "(미지정)"
    base, n = t, 2
    while t in used:
        t, n = f"{base}_{n}", n + 1
    used.add(t)
    return t


def write(money, items, spend, ledger, out, method=None, exclude=()):
    wb = Workbook()
    summary = wb.active
    summary.title = "요약"
    order = list(dict.fromkeys([r["family"] or "(미지정)" for r in money] + [r["family"] or "(미지정)" for r in items]))
    by_money, by_items = defaultdict(list), defaultdict(list)
    for r in money:
        by_money[r["family"] or "(미지정)"].append(r)
    for r in items:
        by_items[r["family"] or "(미지정)"].append(r)
    total = sum(r["amount"] or 0 for r in money)
    spent = sum(r["amount"] or 0 for r in spend)

    # 부의금 입력
    ws = wb.create_sheet("부의금 입력")
    head(ws, 1, 1, ["받은 가족", "이름", "금액", "소속·관계", "메모", "확인 필요", "일차", "원본 위치"])
    for i, r in enumerate(money, 2):
        for j, v in enumerate([r["family"], r["giver"], r["amount"], r["org"], r["memo"], r["check"], r["day"], r["src"]], 1):
            c = ws.cell(i, j, v if v != "" else None)
            c.border = BOX
            if r["check"]:
                c.fill = YELLOW
        ws.cell(i, 3).number_format = WON
        ws.cell(i, 7).number_format = '0"일차"'
    widths(ws, {"A": 12, "B": 16, "C": 13, "D": 20, "E": 18, "F": 22, "G": 8, "H": 34})
    ws.freeze_panes = "A2"

    # 근조 화환
    ws = wb.create_sheet("근조 화환")
    head(ws, 1, 1, ["받은 가족", "보낸 분", "종류", "소속·관계", "메모", "확인 필요", "원본 위치"])
    for i, r in enumerate(items, 2):
        for j, v in enumerate([r["family"], r["giver"], r["kind"], r["org"], r["memo"], r["check"], r["src"]], 1):
            c = ws.cell(i, j, v if v != "" else None)
            c.border = BOX
            if r["check"]:
                c.fill = YELLOW
    widths(ws, {"A": 12, "B": 40, "C": 13, "D": 18, "E": 16, "F": 20, "G": 34})
    ws.freeze_panes = "A2"

    # 지출
    ws = wb.create_sheet("지출")
    head(ws, 1, 1, ["항목", "금액", "특이사항", "낸 사람", "확인 필요"])
    for i, r in enumerate(spend, 2):
        for j, v in enumerate([r["item"], r["amount"], r["memo"], r["payer"], r["check"]], 1):
            c = ws.cell(i, j, v if v != "" else None)
            c.border = BOX
            if r["check"]:
                c.fill = YELLOW
        ws.cell(i, 2).number_format = WON
    ws.cell(1, 7, "지출 정산").font = Font(bold=True, size=14)
    kpi(ws, 3, 7, "부의금 총액", total, WON)
    kpi(ws, 4, 7, "지출 총액", spent, WON)
    kpi(ws, 5, 7, "남는 돈", total - spent, WON)
    head(ws, 8, 7, ["낸 사람", "낸 금액", "건수"])
    payers = defaultdict(lambda: [0, 0])
    for r in spend:
        p = payers[r["payer"] or "(미기재)"]
        p[0] += r["amount"] or 0
        p[1] += 1
    for i, (name, (amt, n)) in enumerate(sorted(payers.items(), key=lambda x: -x[1][0]), 9):
        ws.cell(i, 7, name).border = BOX
        c = ws.cell(i, 8, amt)
        c.number_format, c.border = WON, BOX
        ws.cell(i, 9, n).border = BOX
    if "공금" in payers:
        ws.cell(1, 13, "공금").font = Font(bold=True, size=14)
        kpi(ws, 3, 13, "공금 시작", None, WON)
        kpi(ws, 4, 13, "공금에서 나간 돈", payers["공금"][0], WON)
        ws.cell(3, 15, "← 시작 금액을 알려 주시면 잔액을 계산합니다").font = Font(italic=True, color="999999")
    widths(ws, {"A": 22, "B": 13, "C": 18, "D": 12, "E": 18, "F": 3, "G": 14, "H": 14, "I": 7, "M": 16, "N": 14})
    ws.freeze_panes = "A2"

    # 요약 - 구글 시트 방식과 같은 자리(5행 머리글 '받은 가족', 6행부터 가족)
    s = summary
    kpi(s, 1, 1, "총 합계", total, WON)
    kpi(s, 2, 1, "총 건수", len(money), '0"건"')
    kpi(s, 1, 4, "조화·물품", len(items), '0"개"')
    kpi(s, 2, 4, "확인 필요", sum(1 for r in money + items + spend if r["check"]), '0"건"')
    kpi(s, 1, 8, "지출 총액", spent, WON)
    kpi(s, 2, 8, "남는 돈", total - spent, WON)
    head(s, 5, 1, ["받은 가족", "건수", "합계", "비중", "조화·물품"])
    for i, fam in enumerate(order, 6):
        amt = sum(r["amount"] or 0 for r in by_money[fam])
        vals = [fam, len(by_money[fam]), amt, amt / total if total else 0, len(by_items[fam])]
        for j, v in enumerate(vals, 1):
            s.cell(i, j, v).border = BOX
        s.cell(i, 2).number_format = '0"건"'
        s.cell(i, 3).number_format = WON
        s.cell(i, 4).number_format = "0.0%"
        s.cell(i, 5).number_format = '0"개"'
    # 정산: 사용자가 고른 방법으로만 채운다. '(미지정)' 은 나누지 않는다.
    if method:
        fams = [(f, sum(r["amount"] or 0 for r in by_money[f]), len(by_money[f])) for f in order]
        res = shares(fams, spent, method, set(exclude) | {"(미지정)"})
        head(s, 5, 7, [f"지출 분담 ({METHODS[method]})", "남는 몫", "실제 지급", "절삭된 금액"])
        sums = [0, 0, 0, 0]
        for i, fam in enumerate(order, 6):
            share, left = res[fam]
            left = round(left)
            paid = int(left / 10000) * 10000          # 만 원 단위로 내림(0 쪽으로), 천 원 단위 절삭
            for k, (j, v) in enumerate(((7, round(share)), (8, left), (9, paid), (10, left - paid))):
                c = s.cell(i, j, v)
                c.number_format, c.border = WON, BOX
                sums[k] += v
        t = 6 + len(order)
        s.cell(t, 1, "합계").font = Font(bold=True)
        for k, j in enumerate((7, 8, 9, 10)):
            c = s.cell(t, j, sums[k])
            c.number_format, c.border, c.font, c.fill = WON, BOX, Font(bold=True), BLUE
        note = f"정산 방법: {method}. {METHODS[method]} · 실제 지급은 남는 몫을 만 원 단위로 내린 금액"
        if exclude:
            note += f" · 나눠 내지 않음: {', '.join(exclude)}"
        s.cell(t + 1, 7, note).font = Font(color="777777")

    days = sorted({r["day"] for r in money if r["day"]})
    if days:
        s.cell(4, 12, "일차별 소계").font = Font(bold=True, size=12)
        head(s, 5, 12, ["일차", "건수", "합계", "장부에 적은 소계", "차이"])
        for i, d in enumerate(days, 6):
            amt = sum(r["amount"] or 0 for r in money if r["day"] == d)
            led = ledger.get(d)
            for j, v in enumerate([d, sum(1 for r in money if r["day"] == d), amt, led, (amt - led) if led is not None else None], 12):
                s.cell(i, j, v).border = BOX
            s.cell(i, 12).number_format = '0"일차"'
            for j in (14, 15, 16):
                s.cell(i, j).number_format = WON
    widths(s, {"A": 14, "B": 12, "C": 14, "D": 12, "E": 14, "F": 3, "G": 22, "H": 14, "I": 14, "J": 13,
               "K": 3, "L": 8, "M": 8, "N": 14, "O": 16, "P": 12})
    s.sheet_view.showGridLines = False

    # 물품 통계
    ws = wb.create_sheet("물품 통계")
    ws.cell(1, 1, "조화·물품 통계").font = Font(bold=True, size=14)
    kinds = Counter(r["kind"] or "(종류 없음)" for r in items)
    kind_list = [k for k, _ in kinds.most_common()]
    head(ws, 4, 1, ["종류", "개수", "비중"])
    for i, k in enumerate(kind_list, 5):
        ws.cell(i, 1, k)
        ws.cell(i, 2, kinds[k])
        ws.cell(i, 3, kinds[k] / len(items)).number_format = "0.0%"
    head(ws, 4, 5, ["받은 가족"] + kind_list + ["합계"])
    fams = sorted({r["family"] or "(미지정)" for r in items}, key=lambda f: -len(by_items[f]))
    for i, fam in enumerate(fams, 5):
        ws.cell(i, 5, fam)
        for j, k in enumerate(kind_list, 6):
            ws.cell(i, j, sum(1 for r in by_items[fam] if (r["kind"] or "(종류 없음)") == k) or None)
        ws.cell(i, 6 + len(kind_list), len(by_items[fam])).font = Font(bold=True)
    widths(ws, {"A": 14, "B": 8, "C": 8, "D": 3, "E": 14})
    ws.sheet_view.showGridLines = False

    # 확인 필요
    ws = wb.create_sheet("확인 필요")
    head(ws, 1, 1, ["어디", "받은 가족", "이름·항목", "원래 적힌 값", "이유", "원본 위치"])
    r_ = 2
    for where, rows, name_key in (("부의금", money, "giver"), ("조화·물품", items, "giver"), ("지출", spend, "item")):
        for r in rows:
            if r["check"]:
                for j, v in enumerate([where, r.get("family", ""), r[name_key], r.get("raw", ""), r["check"], r["src"]], 1):
                    ws.cell(r_, j, v)
                r_ += 1
    if r_ == 2:
        ws.cell(2, 1, "확인할 줄이 없습니다.")
    widths(ws, {"A": 10, "B": 12, "C": 18, "D": 14, "E": 26, "F": 34})

    # 사람별 탭 - A1 '받은 가족', B1 이름 (deliver.py 가 이걸로 찾는다)
    used = set(wb.sheetnames)
    for fam in order:
        if fam == "(미지정)":
            continue
        ws = wb.create_sheet(safe_title(fam, used))
        rows_m, rows_i = by_money[fam], by_items[fam]
        amt = sum(r["amount"] or 0 for r in rows_m)
        for c, v in [(1, "받은 가족"), (2, fam), (3, "합계"), (4, amt), (5, "건수"), (6, len(rows_m)),
                     (8, "조화·물품"), (9, len(rows_i))]:
            cell = ws.cell(1, c, v)
            cell.font = Font(bold=True, size=12)
            cell.alignment = Alignment(horizontal="left")
        ws.cell(1, 2).fill = YELLOW
        ws.cell(1, 4).number_format = WON
        ws.cell(1, 6).number_format = '0"건"'
        ws.cell(1, 9).number_format = '0"개"'
        head(ws, 3, 1, ["번호", "이름", "금액", "소속·관계", "메모"])
        head(ws, 3, 8, ["보낸 분", "종류", "소속·관계", "메모"])
        for i, r in enumerate(rows_m, 1):
            for j, v in enumerate([i, r["giver"], r["amount"], r["org"] or None, r["memo"] or None], 1):
                ws.cell(3 + i, j, v).border = BOX
            ws.cell(3 + i, 3).number_format = WON
        for i, r in enumerate(rows_i):
            for j, v in enumerate([r["giver"], r["kind"], r["org"] or None, r["memo"] or None], 8):
                ws.cell(4 + i, j, v).border = BOX
        widths(ws, {"A": 7, "B": 16, "C": 13, "D": 20, "E": 16, "F": 3, "G": 3, "H": 38, "I": 12, "J": 16, "K": 14})
        ws.freeze_panes = "A4"
        ws.sheet_view.showGridLines = False
    wb.save(out)
    return total, spent


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("inputs", nargs="+")
    ap.add_argument("-o", "--out", default="부의금정산.xlsx")
    ap.add_argument("--split", type=int, choices=[1, 2, 3], help="사용자가 고른 정산 방법")
    ap.add_argument("--exclude", default="", help="비용을 나눠 내지 않을 가족, 쉼표로")
    a = ap.parse_args()
    paths = []
    for p in map(lambda x: Path(x).expanduser(), a.inputs):
        paths += sorted(x for x in p.glob("*.xlsx") if not x.name.startswith("~$")) if p.is_dir() else [p]
    money, items, spend, ledger = read_inputs(paths)
    checks(money, items, spend)
    exclude = [x.strip() for x in a.exclude.split(",") if x.strip()]
    total, spent = write(money, items, spend, ledger, a.out, a.split, exclude)
    print(f"부의금 {len(money)}건 {total:,}원 · 조화·물품 {len(items)}개 · 지출 {len(spend)}건 {spent:,}원")
    print(f"확인 필요 {sum(1 for r in money + items + spend if r['check'])}건")
    if ledger:
        print("장부 소계:", ", ".join(f"{d}일차 {v:,}원" for d, v in sorted(ledger.items())))


if __name__ == "__main__":
    main()
