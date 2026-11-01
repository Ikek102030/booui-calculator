"""정산 방법 세 가지로 미리 계산해 표로 보여준다. 사용자가 고르기 전에 쓴다.

읽는 곳(구글 시트 내보내기, build_excel.py 결과 모두 같은 자리):
  요약 탭 '받은 가족' 머리글 아래 A 이름 · B 건수(손님 수) · C 합계(받은 부의금), 지출 탭 H4 지출 총액

    python3 split_preview.py 정산.xlsx [--exclude 이름,이름] [--spend 금액]

방법
  1. 똑같이 나누기     지출 ÷ 참여 가족 수
  2. 받은 부의금 비율   지출 × 내 부의금 ÷ 참여 가족 부의금 합
  3. 손님 수 비율       지출 × 내 손님 수 ÷ 참여 가족 손님 수 합
"""
import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

METHODS = {1: "똑같이 나누기", 2: "받은 부의금 비율", 3: "손님 수 비율"}
NOT_A_FAMILY = {"(미지정)"}  # 받은 가족이 빈 기록 묶음. 사람이 아니라 나누는 대상에서 항상 뺀다
PROS_CONS = {
    1: ("가장 단순하고 설명하기 쉽다", "손님이 적게 온 가족도 똑같이 내서 부담이 클 수 있다"),
    2: ("모두 같은 비율만큼 남아 가장 무난하다", "부의금을 많이 받은 가족이 비용도 많이 낸다"),
    3: ("식대처럼 손님 수만큼 드는 비용을 반영한다", "봉투 하나에 여럿이 와도 1명으로 세고, 금액 차이는 반영하지 않는다"),
}


def num(v):
    return v if isinstance(v, (int, float)) else 0


def read(path):
    wb = load_workbook(path, data_only=True)
    ws = wb["요약"]
    head = next(r for r in range(1, 30) if ws.cell(r, 1).value == "받은 가족")
    fams = []
    r = head + 1
    while ws.cell(r, 1).value not in (None, "", "합계"):
        fams.append((str(ws.cell(r, 1).value), num(ws.cell(r, 3).value), num(ws.cell(r, 2).value)))
        r += 1
    spend = num(wb["지출"]["H4"].value) if "지출" in wb.sheetnames else 0
    return fams, spend


def shares(fams, spend, method, exclude=()):
    """{이름: (지출 분담, 남는 몫)}. 참여하지 않는 가족은 분담 0."""
    exclude = set(exclude) | NOT_A_FAMILY
    weight = {}
    for name, amount, guests in fams:
        if name in exclude:
            weight[name] = 0
        else:
            weight[name] = {1: 1, 2: amount, 3: guests}[method]
    total = sum(weight.values())
    out = {}
    for name, amount, _ in fams:
        share = spend * weight[name] / total if total else 0
        out[name] = (share, amount - share)
    return out


def table(fams, spend, exclude=()):
    res = {m: shares(fams, spend, m, exclude) for m in METHODS}
    lines = ["| 받은 가족 | 받은 부의금 | 손님 | " + " | ".join(f"{m}. {METHODS[m]} 분담 / 남는 몫" for m in METHODS) + " |",
             "|---|---:|---:|" + "---:|" * len(METHODS)]
    for name, amount, guests in fams:
        cells = [f"{res[m][name][0]:,.0f} / {res[m][name][1]:,.0f}" for m in METHODS]
        mark = " (나누지 않음)" if name in NOT_A_FAMILY else " (빠짐)" if name in exclude else ""
        lines.append(f"| {name}{mark} | {amount:,.0f} | {guests:,.0f} | " + " | ".join(cells) + " |")
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--exclude", default="", help="비용을 나눠 내지 않을 가족, 쉼표로")
    ap.add_argument("--spend", type=float, help="지출 총액을 직접 넣을 때")
    a = ap.parse_args()
    fams, spend = read(Path(a.xlsx).expanduser())
    if a.spend is not None:
        spend = a.spend
    exclude = {x.strip() for x in a.exclude.split(",") if x.strip()}
    if not fams:
        sys.exit("요약 탭에서 가족을 찾지 못했습니다.")
    print(f"지출 총액 {spend:,.0f}원 · 참여 {sum(1 for f in fams if f[0] not in exclude | NOT_A_FAMILY)}가족\n")
    print(table(fams, spend, exclude))
    print()
    for m, (pro, con) in PROS_CONS.items():
        print(f"{m}. {METHODS[m]}: 장점 {pro}. 단점 {con}.")


if __name__ == "__main__":
    main()
