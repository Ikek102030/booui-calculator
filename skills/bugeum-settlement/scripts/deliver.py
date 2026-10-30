"""구글 시트에서 내려받은 엑셀 하나 → 가족 한 분당 폴더 [엑셀 1개 + PDF 리포트 1개].

엑셀: 공통 탭(요약·물품 통계·지출·부의금 입력·근조 화환 …)은 그대로 두고,
      '가족별'(드롭다운) 탭 자리에 그분 탭 하나만 넣는다. 다른 가족 탭은 뺀다.
      받는 쪽 엑셀에서 구글 전용 수식이 깨지므로 모든 칸을 계산된 값으로 바꿔 저장한다.
PDF : 요약 카드(받은 부의금·조화·물품·정산 열) + 부의금 명단 + 조화·물품 목록.

사람 탭은 A1 이 '받은 가족' 이고 B1 이 탭 이름과 같은 탭이다(SKILL.md 의 '사람별 탭').

    python3 deliver.py export.xlsx 출력폴더
"""
import datetime
import html
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

INPUT, ITEMS, SUMMARY, PICKER = "부의금 입력", "근조 화환", "요약", "가족별"
DROP = [PICKER, "확인 필요"]  # 받는 분 파일에서 빼는 탭(드롭다운 화면, 점검용)
CHROME_PATHS = [
    os.environ.get("CHROME", ""),
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Chromium.app/Contents/MacOS/Chromium",
    shutil.which("google-chrome") or "",
    shutil.which("chromium") or "",
]


def is_formula(v):
    return isinstance(v, ArrayFormula) or (isinstance(v, str) and v.startswith("="))


def won(n):
    return f"{n:,.0f}원"


def label_of(cell):
    """머리글 칸의 보이는 글자. 값이 0.77 이고 서식이 '"지출 분담 "0%' 이면 '지출 분담 77%'."""
    v = cell.value
    if isinstance(v, str):
        return v
    if isinstance(v, (int, float)):
        text = "".join(re.findall(r'"([^"]*)"', cell.number_format or ""))
        num = f"{v:.0%}" if "%" in (cell.number_format or "") else f"{v:g}"
        return f"{text}{num}".strip()
    return ""


def person_tabs(wb):
    return [ws.title for ws in wb.worksheets
            if str(ws["A1"].value or "").strip() == "받은 가족" and str(ws["B1"].value or "").strip() == ws.title]


def read_people(vals):
    """부의금 입력·근조 화환 탭에서 사람별 명단. '소계'·'합계' 줄은 구분선이라 뺀다."""
    people = {}
    for fam, name, amt, org, memo in vals[INPUT].iter_rows(min_row=2, max_col=5, values_only=True):
        name = str(name or "")
        if not fam or re.search("소계|합계", name):
            continue
        people.setdefault(fam, {"money": [], "items": []})["money"].append(
            (name, amt if isinstance(amt, (int, float)) else None, org or "", memo or ""))
    if ITEMS in vals.sheetnames:
        for fam, who, kind, org, memo in vals[ITEMS].iter_rows(min_row=2, max_col=5, values_only=True):
            if fam and (who or kind):
                people.setdefault(fam, {"money": [], "items": []})["items"].append(
                    (str(who or "").strip(), kind or "", org or "", memo or ""))
    return people


def read_settlement(vals):
    """요약 탭 가족 표에서 '합계' 와, 그 오른쪽 정산 열(지출 분담·남는 몫·실제 지급 …)을 읽는다."""
    ws = vals[SUMMARY]
    head_row = next((r for r in range(1, 30) if ws.cell(r, 1).value == "받은 가족"), None)
    if not head_row:
        return {}, []
    heads = {c: label_of(ws.cell(head_row, c)) for c in range(1, ws.max_column + 1)}
    total_col = next(c for c, h in heads.items() if h == "합계")
    # 정산 열: 가족 표 오른쪽 빈 열(여백) 다음부터 다음 빈 열 전까지
    gap = next((c for c in range(total_col, ws.max_column + 1) if not heads.get(c)), None)
    settle_cols = []
    if gap:
        c = gap + 1
        while heads.get(c):
            settle_cols.append(c)
            c += 1
    out = {}
    r = head_row + 1
    while ws.cell(r, 1).value:
        out[ws.cell(r, 1).value] = {"total": ws.cell(r, total_col).value,
                                    "settle": [(heads[c], ws.cell(r, c).value) for c in settle_cols]}
        r += 1
    return out, [heads[c] for c in settle_cols]


CSS = """
@page { size: A4; margin: 16mm 14mm; }
* { box-sizing: border-box; }
body { font-family: 'Pretendard', 'Apple SD Gothic Neo', 'Malgun Gothic', sans-serif; color: #222; font-size: 10.5pt; margin: 0; }
.kicker { color: #888; font-size: 10pt; }
h1 { font-size: 22pt; margin: 2px 0 14px; font-weight: 700; }
h2 { font-size: 12.5pt; margin: 22px 0 8px; font-weight: 700; }
.cards { display: flex; gap: 8px; flex-wrap: wrap; }
.card { flex: 1 1 0; border: 1px solid #e3e3e3; border-radius: 8px; padding: 10px 12px; min-width: 0; }
.card .l { color: #888; font-size: 9pt; }
.card .v { font-size: 13pt; font-weight: 700; margin-top: 2px; white-space: nowrap; }
.card .s { color: #888; font-size: 9pt; margin-top: 1px; }
.card.hi { background: #f3f7ff; border-color: #cfdcf7; }
table { width: 100%; border-collapse: collapse; }
th { background: #f4f4f4; font-weight: 600; font-size: 9.5pt; color: #555; }
th, td { border-bottom: 1px solid #eaeaea; padding: 5px 8px; text-align: left; }
td.n, th.n { text-align: right; white-space: nowrap; width: 110px; padding-right: 22px; }
td.nm { width: 180px; }
td.i { color: #aaa; width: 34px; text-align: right; }
tr.sum td { font-weight: 700; border-top: 1.5px solid #bbb; background: #fafafa; }
thead { display: table-header-group; }
tr { page-break-inside: avoid; }
.foot { margin-top: 24px; color: #aaa; font-size: 8.5pt; text-align: right; }
"""


def report_html(fam, p, s):
    e = html.escape
    money, items = p["money"], p["items"]
    total = sum(a or 0 for _, a, _, _ in money)
    cards = [f'<div class="card"><div class="l">받은 부의금</div><div class="v">{won(total)}</div><div class="s">{len(money)}건</div></div>',
             f'<div class="card"><div class="l">조화·물품</div><div class="v">{len(items)}개</div></div>']
    if s and s["total"]:
        settle = [(h, v) for h, v in s["settle"] if isinstance(v, (int, float))]
        for i, (h, v) in enumerate(settle):
            hi = " hi" if i == len(settle) - 1 else ""
            cards.append(f'<div class="card{hi}"><div class="l">{e(h)}</div><div class="v">{won(v)}</div></div>')
    out = [f"<!doctype html><html lang='ko'><meta charset='utf-8'><style>{CSS}</style><body>",
           '<div class="kicker">부의금 · 조화 정리</div>', f"<h1>{e(fam)} 님</h1>",
           f'<div class="cards">{"".join(cards)}</div>']
    if money:
        show_org = any(o for _, _, o, _ in money)
        show_memo = any(m for _, _, _, m in money)
        out.append(f"<h2>부의금 명단 <span class='kicker'>{len(money)}건</span></h2><table><thead><tr><th></th><th>이름</th><th class='n'>금액</th>"
                   + ("<th>소속·관계</th>" if show_org else "") + ("<th>메모</th>" if show_memo else "") + "</tr></thead><tbody>")
        for i, (name, amt, org, memo) in enumerate(money, 1):
            out.append(f"<tr><td class='i'>{i}</td><td class='nm'>{e(name)}</td><td class='n'>{won(amt) if amt is not None else '-'}</td>"
                       + (f"<td>{e(str(org))}</td>" if show_org else "") + (f"<td>{e(str(memo))}</td>" if show_memo else "") + "</tr>")
        out.append(f"<tr class='sum'><td></td><td>합계</td><td class='n'>{won(total)}</td>" + "<td></td>" * (show_org + show_memo) + "</tr></tbody></table>")
    if items:
        out.append(f"<h2>조화·물품 <span class='kicker'>{len(items)}개</span></h2><table><thead><tr><th></th><th>보낸 분</th><th>종류</th></tr></thead><tbody>")
        for i, (who, kind, _, _) in enumerate(items, 1):
            out.append(f"<tr><td class='i'>{i}</td><td>{e(who)}</td><td>{e(str(kind))}</td></tr>")
        out.append("</tbody></table>")
    today = datetime.date.today()
    out.append(f'<div class="foot">{today.year}년 {today.month}월 {today.day}일 정리</div></body></html>')
    return "".join(out), total


def person_workbook(export, person, people, out):
    wb = load_workbook(export)
    vals = load_workbook(export, data_only=True)
    for name in DROP + [p for p in people if p != person]:
        if name in wb.sheetnames:
            del wb[name]
    for ws in wb.worksheets:
        src = vals[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if is_formula(c.value):  # 배열 수식(ArrayFormula)도 빠뜨리지 않는다
                    c.value = src[c.coordinate].value
    me = wb[person]
    me.sheet_state = "visible"
    rest = [ws for ws in wb.worksheets if ws.title not in (SUMMARY, person)]
    wb._sheets = [wb[SUMMARY], me] + rest
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == SUMMARY
    wb.save(out)
    return wb.sheetnames


def main():
    export, dest = Path(sys.argv[1]).expanduser(), Path(sys.argv[2]).expanduser()
    vals = load_workbook(export, data_only=True)
    tabs = person_tabs(vals)
    if not tabs:
        sys.exit("사람별 탭이 없습니다. SKILL.md 의 '사람별 탭' 단계를 먼저 하세요.")
    people = read_people(vals)
    settle, _ = read_settlement(vals)
    chrome = next((c for c in CHROME_PATHS if c and Path(c).exists()), None)
    if not chrome:
        print("크롬을 찾지 못해 PDF 는 건너뜁니다. CHROME 환경변수로 경로를 알려 주세요.")
    tmp = Path(tempfile.mkdtemp())
    grand = 0
    for person in tabs:
        d = dest / person
        d.mkdir(parents=True, exist_ok=True)
        xlsx = d / f"부의금 계산_{person}.xlsx"
        sheets = person_workbook(export, person, tabs, xlsx)

        # 검산: 그분 탭 명단 합계 = 그분 탭 D1 = 요약 합계 = PDF 합계, 수식 0개
        chk = load_workbook(xlsx)
        ws = chk[person]
        # 번호(A열)가 있는 줄은 이름이 비어도 센다
        listed = sum(ws.cell(r, 3).value for r in range(4, ws.max_row + 1)
                     if isinstance(ws.cell(r, 1).value, (int, float)) and isinstance(ws.cell(r, 3).value, (int, float)))
        assert listed == (ws["D1"].value or 0), (person, listed, ws["D1"].value)
        assert not any(is_formula(c.value) for s in chk.worksheets for r in s.iter_rows() for c in r), person
        s = settle.get(person)
        if s and s["total"] is not None:
            assert listed == s["total"], (person, listed, s["total"])
        doc, total = report_html(person, people.get(person, {"money": [], "items": []}), s)
        assert total == listed, (person, total, listed)

        if chrome:
            h = tmp / f"{person}.html"
            h.write_text(doc, encoding="utf-8")
            subprocess.run([chrome, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                            f"--print-to-pdf={d / f'부의금 정리_{person}.pdf'}", h.as_uri()],
                           check=True, capture_output=True)
        grand += listed
        print(f"{person}: {listed:,.0f}원 · 탭 {sheets}")
    print(f"합계 {grand:,.0f}원 · {len(tabs)}명")


if __name__ == "__main__":
    main()
