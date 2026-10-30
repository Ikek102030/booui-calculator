"""AI 없이 쓰는 부의금 정산 엑셀 양식을 만든다.

엑셀 2010 이후 · 구글 시트 · 한셀에서 모두 돌도록 기본 함수(SUMIF·COUNTIF·INDEX/MATCH)만 쓴다.
FILTER·UNIQUE·LET 같은 최신 함수는 옛 엑셀에서 깨지므로 쓰지 않는다.
대신 가족 목록은 사람이 '가족 목록' 탭에 적고, 가족별 명단은 숨긴 순번 열로 INDEX/MATCH 한다.

    python3 make_template.py 부의금_정산_양식.xlsx
"""
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from functools import partial

from openpyxl.styles import Alignment, Border, PatternFill, Side
from openpyxl.styles import Font as _Font

Font = partial(_Font, name="맑은 고딕")
from openpyxl.worksheet.datavalidation import DataValidation

N_MONEY, N_ITEM, N_SPEND, N_FAM, N_DAY, N_VIEW = 1000, 500, 300, 20, 10, 400
M, I, S = N_MONEY + 1, N_ITEM + 1, N_SPEND + 1  # 마지막 행 번호
F_LAST = N_FAM + 1
KINDS = ["근조화환", "근조 바구니", "근조기", "쌀화환", "기타"]

WON = '#,##0"원"'
GREY = PatternFill("solid", fgColor="EDEDED")
INPUT = PatternFill("solid", fgColor="FFF8E1")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
BLUE = PatternFill("solid", fgColor="DCE6F7")
LINE = Side(style="thin", color="C8C8C8")  # 배경색을 칠하면 엑셀 눈금선이 가려지므로 칸마다 테두리를 그린다
BOX = Border(top=LINE, bottom=LINE, left=LINE, right=LINE)
BOLD = Font(bold=True)
MUTED = Font(color="777777")
ORANGE, SKY = "F4B342", "4285F4"

MONEY = "'부의금 입력'"
ITEM = "'근조 화환'"
FAMS = f"'가족 목록'!$A$2:$A${F_LAST}"


def head(ws, row, col, labels, fill=GREY):
    for i, h in enumerate(labels):
        c = ws.cell(row, col + i, h)
        c.font, c.fill, c.border = BOLD, fill, BOX
        c.alignment = Alignment(horizontal="center", vertical="center")


def widths(ws, d):
    for col, w in d.items():
        ws.column_dimensions[col].width = w


def label(ws, ref, text, big=False):
    ws[ref] = text
    ws[ref].font = Font(bold=True, size=14) if big else MUTED


def value(ws, ref, formula, fmt=WON, size=13):
    ws[ref] = formula
    ws[ref].font = Font(bold=True, size=size)
    ws[ref].number_format = fmt


def list_validation(ws, sqref, source, strict=False):
    dv = DataValidation(type="list", formula1=source, allow_blank=True, showDropDown=False)
    dv.showErrorMessage = strict  # False 면 목록 밖 값도 받는다
    ws.add_data_validation(dv)
    dv.add(sqref)


def is_sep(ref):
    """'소계'·'합계' 줄(구분선)인지."""
    return f'OR(ISNUMBER(SEARCH("소계",{ref})),ISNUMBER(SEARCH("합계",{ref})))'


def guide(wb):
    ws = wb.active
    ws.title = "사용법"
    lines = [
        ("부의금 정산 양식", Font(bold=True, size=18)),
        ("", None),
        ("삼가 고인의 명복을 빕니다.", Font(size=12)),
        ("경황 없으실 텐데, 적기만 하시면 계산은 이 파일이 합니다.", Font(size=11, color="555555")),
        ("", None),
        ("① '가족 목록' 탭에 부의금을 받는 가족 이름을 적으세요. (최대 20명)", None),
        ("    여럿이 함께 받은 돈은 '공동'이나 '아빠·작은아빠'처럼 이름을 하나 만들어 쓰세요.", MUTED),
        ("② '부의금 입력' 탭에 한 줄에 한 분씩 적으세요. 받은 가족은 목록에서 고르고, 금액은 숫자로(50000).", None),
        ("③ 하루가 끝나면 이름 칸에 '1일차 소계', 금액 칸에 그날 합계를 적으세요. 요약 탭에서 실제 합계와 비교해 줍니다.", None),
        ("    마지막에 '총합계' 줄을 적어도 됩니다. '소계'·'합계' 줄은 합계에 두 번 들어가지 않습니다.", MUTED),
        ("④ 화환·근조기·바구니는 '근조 화환' 탭, 장례 비용은 '지출' 탭에 적으세요.", None),
        ("    부의금과 따로 모은 공금이 있으면 지출 탭 오른쪽 '공금 시작'에 금액을 넣고, 낸 사람을 '공금'으로 고르세요.", MUTED),
        ("⑤ '요약' 탭에서 가족별 합계·지출 분담·남는 몫을, '가족별' 탭에서 노란 칸의 이름을 골라 그분 명단을 봅니다.", None),
        ("", None),
        ("노란색 줄은 확인이 필요한 줄입니다. (금액이 글자로 적힘, 이름이나 받은 가족이 빔, 같은 줄이 두 번 적힘 등)", None),
        ("같은 분이 근조기와 화환을 함께 보내신 건 중복으로 보지 않습니다.", MUTED),
        ("연한 노란 칸이 적는 칸입니다. 나머지 칸은 수식이니 지우지 마세요.", MUTED),
        ("", None),
        (f"부의금 {N_MONEY}줄, 화환 {N_ITEM}줄, 지출 {N_SPEND}줄까지 적을 수 있습니다. 구글 시트에 올려 써도 됩니다.", MUTED),
    ]
    for i, (t, f) in enumerate(lines, 1):
        ws.cell(i, 2, t)
        if f:
            ws.cell(i, 2).font = f
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    ws.sheet_view.showGridLines = False


def families(wb):
    ws = wb.create_sheet("가족 목록")
    head(ws, 1, 1, ["받은 가족"])
    for r in range(2, F_LAST + 1):
        ws.cell(r, 1).fill = INPUT
        ws.cell(r, 1).border = BOX
    ws["C1"] = "부의금·화환을 받는 가족 이름을 적으세요. 여기 적은 이름이 드롭다운과 요약에 나옵니다."
    ws["C1"].font = MUTED
    widths(ws, {"A": 18, "C": 80})
    ws.sheet_properties.tabColor = ORANGE


def money(wb):
    ws = wb.create_sheet("부의금 입력")
    head(ws, 1, 1, ["받은 가족", "이름", "금액", "소속·관계", "메모"])
    head(ws, 1, 6, ["확인 필요", "일차", "가족별 순번", "소계 번호", "확인 표시"], BLUE)
    for r in range(2, M + 1):
        for c in range(1, 6):
            ws.cell(r, c).fill = INPUT
        for c in range(1, 8):
            ws.cell(r, c).border = BOX
        ws.cell(r, 3).number_format = WON
        ws.cell(r, 6).value = (
            f'=IF(COUNTA(A{r}:E{r})=0,"",IF({is_sep(f"B{r}")},"",TRIM('
            f'IF(B{r}="","이름 없음 ","")&IF(A{r}="","받은 가족 없음 ","")'
            f'&IF(C{r}="","금액 없음 ",IF(ISNUMBER(C{r}),IF(C{r}<1000,"금액 확인 ",""),"금액 확인 "))'
            f'&IF(AND(B{r}<>"",COUNTIFS($A$2:$A${M},A{r},$B$2:$B${M},B{r},$C$2:$C${M},C{r})>1),"중복 의심",""))))')
        ws.cell(r, 7).value = f'=IF(OR(COUNTA(A{r}:E{r})=0,{is_sep(f"B{r}")}),"",1+COUNTIF($B$1:B{r - 1},"*소계*"))'
        ws.cell(r, 7).number_format = '0"일차"'
        ws.cell(r, 8).value = (f'=IF(AND(A{r}<>"",A{r}=가족별!$B$1,G{r}<>""),'
                               f'COUNTIFS($A$2:A{r},가족별!$B$1,$G$2:G{r},">0"),"")')
        ws.cell(r, 9).value = f'=IF(ISNUMBER(SEARCH("소계",B{r})),COUNTIF($B$2:B{r},"*소계*"),"")'
        ws.cell(r, 10).value = f'=IF(F{r}="",0,1)'
    list_validation(ws, f"A2:A{M}", f"={FAMS}")
    ws.conditional_formatting.add(f"A2:G{M}", FormulaRule(formula=[is_sep("$B2")], fill=BLUE, font=BOLD, stopIfTrue=True))
    ws.conditional_formatting.add(f"A2:G{M}", FormulaRule(formula=['$F2<>""'], fill=YELLOW))
    widths(ws, {"A": 13, "B": 16, "C": 13, "D": 20, "E": 18, "F": 22, "G": 8})
    ws.column_dimensions["H"].hidden = True
    ws.column_dimensions["I"].hidden = True
    ws.column_dimensions["J"].hidden = True
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = ORANGE


def items(wb):
    ws = wb.create_sheet("근조 화환")
    head(ws, 1, 1, ["받은 가족", "보낸 분", "종류", "소속·관계", "메모"])
    head(ws, 1, 6, ["확인 필요", "가족별 순번", "기록 있음", "확인 표시"], BLUE)
    for r in range(2, I + 1):
        for c in range(1, 6):
            ws.cell(r, c).fill = INPUT
        for c in range(1, 7):
            ws.cell(r, c).border = BOX
        # 같은 분이 종류를 달리해 보내는 건 정상 — 보낸 분과 종류가 둘 다 같을 때만 의심
        ws.cell(r, 6).value = (
            f'=IF(COUNTA(A{r}:E{r})=0,"",TRIM(IF(A{r}="","받은 가족 없음 ","")&IF(B{r}="","보낸 분 없음 ","")'
            f'&IF(AND(B{r}<>"",COUNTIFS($B$2:$B${I},B{r},$C$2:$C${I},C{r})>1),"중복 의심","")))')
        ws.cell(r, 7).value = f'=IF(AND(A{r}<>"",A{r}=가족별!$B$1),COUNTIF($A$2:A{r},가족별!$B$1),"")'
        ws.cell(r, 8).value = f'=IF(COUNTA(B{r}:C{r})>0,1,0)'
        ws.cell(r, 9).value = f'=IF(F{r}="",0,1)'
    list_validation(ws, f"A2:A{I}", f"={FAMS}")
    list_validation(ws, f"C2:C{I}", '"' + ",".join(KINDS) + '"')
    ws.conditional_formatting.add(f"A2:F{I}", FormulaRule(formula=['$F2<>""'], fill=YELLOW))
    widths(ws, {"A": 13, "B": 40, "C": 13, "D": 18, "E": 16, "F": 20})
    for col in ("G", "H", "I"):
        ws.column_dimensions[col].hidden = True
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = ORANGE


def spend(wb):
    ws = wb.create_sheet("지출")
    head(ws, 1, 1, ["항목", "금액", "특이사항", "낸 사람"])
    head(ws, 1, 5, ["확인 필요"], BLUE)
    for r in range(2, S + 1):
        for c in range(1, 5):
            ws.cell(r, c).fill = INPUT
        for c in range(1, 6):
            ws.cell(r, c).border = BOX
        ws.cell(r, 2).number_format = WON
        ws.cell(r, 5).value = (
            f'=IF(COUNTA(A{r}:D{r})=0,"",IF({is_sep(f"A{r}")},"",TRIM(IF(A{r}="","항목 없음 ","")'
            f'&IF(B{r}="","금액 없음 ",IF(ISNUMBER(B{r}),"","금액 확인 "))&IF(D{r}="","낸 사람 없음",""))))')
        ws.cell(r, 25).value = f'=IF(E{r}="",0,1)'
    label(ws, "G1", "지출 정산", big=True)
    label(ws, "G3", "부의금 총액")
    value(ws, "H3", "=요약!B1")
    label(ws, "G4", "지출 총액")
    value(ws, "H4", f'=SUM(B2:B{S})-SUMIF(A2:A{S},"*합계*",B2:B{S})-SUMIF(A2:A{S},"*소계*",B2:B{S})')
    label(ws, "G5", "남는 돈")
    value(ws, "H5", "=H3-H4")
    ws["G5"].fill = ws["H5"].fill = BLUE
    label(ws, "G7", "공금 시작")
    ws["H7"].fill, ws["H7"].number_format, ws["H7"].border = INPUT, WON, BOX
    ws["I7"] = "← 부의금과 따로 모은 돈이 있으면 적으세요"
    ws["I7"].font = Font(italic=True, color="999999")
    label(ws, "G8", "공금에서 나간 돈")
    value(ws, "H8", f'=SUMIF(D2:D{S},"공금",B2:B{S})')
    label(ws, "G9", "공금 잔액")
    value(ws, "H9", '=IF(H7="","",H7-H8)')
    head(ws, 11, 7, ["낸 사람", "낸 금액", "건수"])
    ws["G12"] = "공금"
    for k in range(N_FAM):
        ws.cell(13 + k, 7).value = f"=IF('가족 목록'!A{2 + k}=\"\",\"\",'가족 목록'!A{2 + k})"
    last = 12 + N_FAM
    for r in range(12, last + 1):
        ws.cell(r, 8).value = f'=IF(G{r}="","",SUMIF($D$2:$D${S},G{r},$B$2:$B${S}))'
        ws.cell(r, 9).value = f'=IF(G{r}="","",COUNTIF($D$2:$D${S},G{r}))'
        ws.cell(r, 8).number_format = WON
    r = last + 1
    ws.cell(r, 7, "(낸 사람 없음)")
    ws.cell(r, 8).value = f'=SUMPRODUCT(($D$2:$D${S}="")*($A$2:$A${S}<>""),$B$2:$B${S})'
    ws.cell(r, 9).value = f'=SUMPRODUCT(($D$2:$D${S}="")*($A$2:$A${S}<>""))'
    ws.cell(r, 8).number_format = WON
    for rr in range(12, r + 1):
        for c in (7, 8, 9):
            ws.cell(rr, c).border = BOX
    list_validation(ws, f"D2:D{S}", f"=$G$12:$G${last}")
    ws.conditional_formatting.add(f"A2:E{S}", FormulaRule(formula=['$E2<>""'], fill=YELLOW))
    widths(ws, {"A": 22, "B": 13, "C": 18, "D": 12, "E": 18, "F": 3, "G": 16, "H": 15, "I": 7})
    ws.column_dimensions["Y"].hidden = True
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = ORANGE


def summary(wb):
    ws = wb.create_sheet("요약", 1)
    label(ws, "A1", "총 합계")
    value(ws, "B1", f"=SUMIF({MONEY}!$G$2:$G${M},\">0\",{MONEY}!$C$2:$C${M})", size=15)
    label(ws, "A2", "총 건수")
    value(ws, "B2", f"=COUNTIF({MONEY}!$G$2:$G${M},\">0\")", '0"건"')
    label(ws, "A3", "확인 필요")
    # 빈 글자("") 세는 방식이 프로그램마다 달라서, 각 탭의 숨긴 '확인 표시'(0/1) 열을 더한다
    value(ws, "B3", f"=SUM({MONEY}!$J$2:$J${M})+SUM({ITEM}!$I$2:$I${I})+SUM(지출!$Y$2:$Y${S})", '0"건"')
    label(ws, "D1", "조화·물품")
    value(ws, "E1", f"=SUM({ITEM}!$H$2:$H${I})", '0"개"')
    label(ws, "D2", "지출 총액")
    value(ws, "E2", "=지출!H4")
    label(ws, "D3", "남는 돈")
    value(ws, "E3", "=지출!H5")
    label(ws, "G1", "장부 총합계")
    value(ws, "H1", f'=IF(COUNTIF({MONEY}!$B$2:$B${M},"*총합계*")=0,"",SUMIF({MONEY}!$B$2:$B${M},"*총합계*",{MONEY}!$C$2:$C${M}))')
    label(ws, "G2", "장부와 차이")
    value(ws, "H2", '=IF(H1="","",B1-H1)')

    head(ws, 5, 1, ["받은 가족", "건수", "합계", "비중", "조화·물품"])
    head(ws, 5, 7, ["지출 분담", "남는 몫", "실제 지급"])
    first, last = 6, 5 + N_FAM
    for k in range(N_FAM):
        r = first + k
        ws.cell(r, 1).value = f"=IF('가족 목록'!A{2 + k}=\"\",\"\",'가족 목록'!A{2 + k})"
        ws.cell(r, 2).value = f'=IF($A{r}="","",COUNTIFS({MONEY}!$A$2:$A${M},$A{r},{MONEY}!$G$2:$G${M},">0"))'
        ws.cell(r, 3).value = f'=IF($A{r}="","",SUMIFS({MONEY}!$C$2:$C${M},{MONEY}!$A$2:$A${M},$A{r},{MONEY}!$G$2:$G${M},">0"))'
        ws.cell(r, 4).value = f'=IF($A{r}="","",IF($B$1=0,0,C{r}/$B$1))'
        ws.cell(r, 5).value = f'=IF($A{r}="","",COUNTIF({ITEM}!$A$2:$A${I},$A{r}))'
        ws.cell(r, 7).value = f'=IF($A{r}="","",IF($B$1=0,0,C{r}*지출!$H$4/$B$1))'
        ws.cell(r, 8).value = f'=IF($A{r}="","",C{r}-G{r})'
        ws.cell(r, 9).fill = INPUT
        for c, f in ((2, '0"건"'), (3, WON), (4, "0.0%"), (5, '0"개"'), (7, WON), (8, WON), (9, WON)):
            ws.cell(r, c).number_format = f
        for c in (1, 2, 3, 4, 5, 7, 8, 9):
            ws.cell(r, c).border = BOX
    t = last + 1
    ws.cell(t, 1, "합계")
    for c, col in ((2, "B"), (3, "C"), (4, "D"), (5, "E"), (7, "G"), (8, "H"), (9, "I")):
        ws.cell(t, c).value = f"=SUM({col}{first}:{col}{last})"
        ws.cell(t, c).number_format = ws.cell(first, c).number_format
    for c in (1, 2, 3, 4, 5, 7, 8, 9):
        ws.cell(t, c).font, ws.cell(t, c).fill, ws.cell(t, c).border = BOLD, BLUE, BOX
    ws.cell(t + 1, 1, "가족 목록에 없는 기록").font = MUTED
    ws.cell(t + 1, 3).value = f"=B1-C{t}"
    ws.cell(t + 1, 3).number_format = WON
    ws.cell(t + 1, 3).font = Font(italic=True, color="993333")
    ws.cell(t + 2, 7, "지출 분담 = 받은 부의금 비율대로 지출을 나눈 금액. 다르게 나누셨다면 '실제 지급'에 적으세요.").font = MUTED

    label(ws, "K4", "일차별 소계")
    ws["K4"].font = Font(bold=True, size=12)
    head(ws, 5, 11, ["일차", "건수", "합계", "장부에 적은 소계", "차이"])
    for d in range(1, N_DAY + 1):
        r = 5 + d
        ws.cell(r, 11).value = f'=IF({d}<=MAX({MONEY}!$G$2:$G${M}),{d},"")'
        ws.cell(r, 12).value = f'=IF(K{r}="","",COUNTIF({MONEY}!$G$2:$G${M},K{r}))'
        ws.cell(r, 13).value = f'=IF(K{r}="","",SUMIF({MONEY}!$G$2:$G${M},K{r},{MONEY}!$C$2:$C${M}))'
        ws.cell(r, 14).value = (f'=IF(K{r}="","",IF(COUNTIF({MONEY}!$I$2:$I${M},K{r})=0,"",'
                                f'SUMIF({MONEY}!$I$2:$I${M},K{r},{MONEY}!$C$2:$C${M})))')
        ws.cell(r, 15).value = f'=IF(OR(K{r}="",N{r}=""),"",M{r}-N{r})'
        for c, f in ((11, '0"일차"'), (12, '0"건"'), (13, WON), (14, WON), (15, WON)):
            ws.cell(r, c).number_format = f
            ws.cell(r, c).border = BOX
    widths(ws, {"A": 16, "B": 14, "C": 15, "D": 11, "E": 14, "F": 3, "G": 16, "H": 15, "I": 14, "J": 3,
                "K": 8, "L": 8, "M": 14, "N": 16, "O": 12})
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = SKY


def by_family(wb):
    ws = wb.create_sheet("가족별", 2)
    for c, v in [(1, "받은 가족"), (3, "합계"), (5, "건수"), (8, "조화·물품")]:
        ws.cell(1, c, v).font = Font(bold=True, size=12)
    ws["B1"] = "='가족 목록'!A2"
    ws["B1"].font, ws["B1"].fill = Font(bold=True, size=12), YELLOW
    list_validation(ws, "B1", f"={FAMS}", strict=True)
    value(ws, "D1", f'=SUMIFS({MONEY}!$C$2:$C${M},{MONEY}!$A$2:$A${M},$B$1,{MONEY}!$G$2:$G${M},">0")', size=12)
    value(ws, "F1", f'=COUNTIFS({MONEY}!$A$2:$A${M},$B$1,{MONEY}!$G$2:$G${M},">0")', '0"건"', size=12)
    value(ws, "I1", f"=COUNTIF({ITEM}!$A$2:$A${I},$B$1)", '0"개"', size=12)
    for c in ("D1", "F1", "I1"):
        ws[c].alignment = Alignment(horizontal="left")
    ws["A2"] = "▲ 노란 칸에서 이름을 고르면 그분 명단이 나옵니다."
    ws["A2"].font = Font(italic=True, color="999999")
    head(ws, 3, 1, ["번호", "이름", "금액", "소속·관계", "메모"])
    head(ws, 3, 8, ["보낸 분", "종류", "소속·관계", "메모"])
    for k in range(1, N_VIEW + 1):
        r = 3 + k
        ws.cell(r, 1).value = f'=IF({k}<=$F$1,{k},"")'
        m = f"MATCH($A{r},{MONEY}!$H$2:$H${M},0)"
        ws.cell(r, 2).value = f'=IF($A{r}="","",INDEX({MONEY}!$B$2:$B${M},{m})&"")'
        ws.cell(r, 3).value = f'=IF($A{r}="","",INDEX({MONEY}!$C$2:$C${M},{m}))'
        ws.cell(r, 4).value = f'=IF($A{r}="","",INDEX({MONEY}!$D$2:$D${M},{m})&"")'
        ws.cell(r, 5).value = f'=IF($A{r}="","",INDEX({MONEY}!$E$2:$E${M},{m})&"")'
        ws.cell(r, 3).number_format = WON
        for c in (1, 2, 3, 4, 5, 8, 9, 10, 11):
            ws.cell(r, c).border = BOX
        mi = f"MATCH({k},{ITEM}!$G$2:$G${I},0)"
        for c, col in ((8, "B"), (9, "C"), (10, "D"), (11, "E")):
            ws.cell(r, c).value = f'=IF({k}>$I$1,"",INDEX({ITEM}!${col}$2:${col}${I},{mi})&"")'
    widths(ws, {"A": 7, "B": 16, "C": 13, "D": 20, "E": 16, "F": 6, "G": 3, "H": 38, "I": 12, "J": 16, "K": 14})
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = SKY


def item_stats(wb):
    ws = wb.create_sheet("물품 통계", 3)
    label(ws, "A1", "조화·물품 통계", big=True)
    head(ws, 4, 1, ["종류", "개수", "비중"])
    for k, kind in enumerate(KINDS + ["그 밖"]):
        r = 5 + k
        ws.cell(r, 1, kind)
        if kind == "그 밖":
            ws.cell(r, 2).value = f"=요약!E1-SUM(B5:B{r - 1})"
        else:
            ws.cell(r, 2).value = f'=COUNTIF({ITEM}!$C$2:$C${I},A{r})'
        ws.cell(r, 3).value = f'=IF(요약!$E$1=0,0,B{r}/요약!$E$1)'
        ws.cell(r, 3).number_format = "0.0%"
    head(ws, 4, 5, ["받은 가족"] + KINDS + ["합계"])
    for k in range(N_FAM):
        r = 5 + k
        ws.cell(r, 5).value = f"=IF('가족 목록'!A{2 + k}=\"\",\"\",'가족 목록'!A{2 + k})"
        for j, kind in enumerate(KINDS):
            ws.cell(r, 6 + j).value = f'=IF($E{r}="","",COUNTIFS({ITEM}!$A$2:$A${I},$E{r},{ITEM}!$C$2:$C${I},"{kind}"))'
        ws.cell(r, 6 + len(KINDS)).value = f'=IF($E{r}="","",COUNTIF({ITEM}!$A$2:$A${I},$E{r}))'
    widths(ws, {"A": 14, "B": 8, "C": 8, "D": 3, "E": 14, "F": 11, "G": 12, "H": 9, "I": 9, "J": 8, "K": 8})
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = SKY


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "부의금_정산_양식.xlsx"
    wb = Workbook()
    # 기본 글꼴: 한국어 엑셀 기본인 맑은 고딕(없으면 각 프로그램 기본 글꼴로 바뀐다)
    wb._fonts[0] = Font(name="맑은 고딕", size=10)
    guide(wb)
    families(wb)
    money(wb)
    items(wb)
    spend(wb)
    summary(wb)
    by_family(wb)
    item_stats(wb)
    # 탭 순서: 사용법 · 요약 · 가족별 · 물품 통계 · 가족 목록 · 부의금 입력 · 근조 화환 · 지출
    order = ["사용법", "요약", "가족별", "물품 통계", "가족 목록", "부의금 입력", "근조 화환", "지출"]
    wb._sheets = [wb[t] for t in order]
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
